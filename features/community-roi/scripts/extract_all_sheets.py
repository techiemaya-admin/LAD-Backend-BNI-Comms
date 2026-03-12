#!/usr/bin/env python3
"""
Extract multiple data types from Excel with proper structure handling
Supports: One-to-One Matrix, Referral Matrix, Combination Matrix, TYFCB Report
Generates SQL to import into database
"""

import openpyxl
import sys
import json

def read_excel_matrix(file_path, sheet_pattern):
    """Read Excel matrix with proper structure (One-to-One or Referral)"""
    try:
        workbook = openpyxl.load_workbook(file_path)
        
        # Find the sheet matching pattern
        matrix_sheet = None
        found_sheet_name = None
        for sheet_name in workbook.sheetnames:
            if sheet_pattern.lower() in sheet_name.lower() and 'matrix' in sheet_name.lower():
                found_sheet_name = sheet_name
                matrix_sheet = workbook[sheet_name]
                break
        
        if not matrix_sheet:
            print(f"⚠️  Could not find {sheet_pattern} Matrix sheet")
            return {}
        
        print(f"📋 Reading {sheet_pattern} from sheet: {found_sheet_name}")
        
        data = {}
        all_rows = list(matrix_sheet.iter_rows(values_only=True))
        
        if len(all_rows) < 3:
            print(f"❌ {sheet_pattern} sheet structure is invalid")
            return {}
        
        # Parse structure:
        # Row 0: Title
        # Row 1: Member | member_b1 | member_b2 | ...
        # Rows 2+: member_a | count | count | ...
        
        header_row = all_rows[1]
        member_b_list = []
        
        # Extract column B headers (member names)
        for i in range(1, len(header_row)):
            member_b = header_row[i]
            if member_b and '__EMPTY' not in str(member_b):
                member_b_list.append((i, str(member_b).strip()))
        
        print(f"  ✅ Found {len(member_b_list)} members across")
        
        # Extract data
        row_count = 0
        for row_data in all_rows[2:]:
            member_a = row_data[0]
            if not member_a or '__EMPTY' in str(member_a) or str(member_a).lower() == 'member':
                continue
            
            member_a = str(member_a).strip()
            row_count += 1
            
            # Get counts for each column member
            for col_idx, member_b in member_b_list:
                try:
                    count_val = row_data[col_idx] if col_idx < len(row_data) else None
                    
                    # Skip zero/empty
                    if count_val is None or str(count_val) in ('0', 'nan', '', None):
                        continue
                    
                    # Skip self
                    if member_a.lower().strip() == member_b.lower().strip():
                        continue
                    
                    count = int(float(count_val))
                    if count > 0:
                        data[f"{member_a}|{member_b}"] = {
                            'a': member_a,
                            'b': member_b,
                            'c': count
                        }
                except (ValueError, TypeError):
                    continue
        
        print(f"  ✅ Processed {row_count} rows")
        record_count = len(data)
        print(f"  ✅ Found {record_count} records with count > 0")
        return data
    
    except Exception as e:
        print(f"❌ Error reading {sheet_pattern}: {e}")
        return {}

def read_excel_heatmap(file_path):
    """Read One-to-One Matrix"""
    return read_excel_matrix(file_path, 'one-to-one')

def read_excel_referral_matrix(file_path):
    """Read Referral Matrix"""
    return read_excel_matrix(file_path, 'referral')

def read_excel_combination_matrix(file_path):
    """Read Combination Matrix (interaction + referral combined)"""
    try:
        workbook = openpyxl.load_workbook(file_path)
        
        # Find Combination Matrix sheet
        comb_sheet = None
        found_sheet_name = None
        for sheet_name in workbook.sheetnames:
            if 'combination' in sheet_name.lower() and 'matrix' in sheet_name.lower():
                found_sheet_name = sheet_name
                comb_sheet = workbook[sheet_name]
                break
        
        if not comb_sheet:
            print("⚠️  Could not find Combination Matrix sheet")
            return {}
        
        print(f"📋 Reading Combination from sheet: {found_sheet_name}")
        
        data = {}
        all_rows = list(comb_sheet.iter_rows(values_only=True))
        
        if len(all_rows) < 3:
            print("❌ Combination sheet structure is invalid")
            return {}
        
        # Combination Matrix: Member A | Member B | Type (M/R/MR)
        # Row 0: Title
        # Row 1: Headers
        # Rows 2+: Data
        
        row_count = 0
        for row_data in all_rows[2:]:
            if not row_data or not row_data[0]:
                continue
            
            member_a = str(row_data[0]).strip()
            if '__EMPTY' in member_a or member_a.lower() == 'member':
                continue
            
            member_b_val = row_data[1] if len(row_data) > 1 else None
            if not member_b_val:
                continue
            
            member_b = str(member_b_val).strip()
            if '__EMPTY' in member_b:
                continue
            
            combo_type_val = row_data[2] if len(row_data) > 2 else None
            if not combo_type_val or member_a.lower() == member_b.lower():
                continue
            
            combo_type = str(combo_type_val).strip().upper()
            
            # Type mapping: M=Meeting, R=Referral, MR=Both
            type_map = {'M': 1, 'R': 2, 'MR': 3, 'RM': 3}
            type_code = type_map.get(combo_type, 0)
            
            if type_code > 0:
                data[f"{member_a}|{member_b}"] = {
                    'a': member_a,
                    'b': member_b,
                    't': type_code
                }
                row_count += 1
        
        print(f"  ✅ Processed {row_count} rows")
        record_count = len(data)
        print(f"  ✅ Found {record_count} combination records")
        return data
    
    except Exception as e:
        print(f"❌ Error reading Combination Matrix: {e}")
        return {}

def read_excel_tyfcb_report(file_path):
    """Read TYFCB Report (contribution scores)"""
    try:
        workbook = openpyxl.load_workbook(file_path)
        
        # Find TYFCB or TYDCB sheet
        tyfcb_sheet = None
        found_sheet_name = None
        for sheet_name in workbook.sheetnames:
            lower_name = sheet_name.lower()
            if ('tyfcb' in lower_name or 'tydcb' in lower_name or 'contribution' in lower_name):
                found_sheet_name = sheet_name
                tyfcb_sheet = workbook[sheet_name]
                break
        
        if not tyfcb_sheet:
            print("⚠️  Could not find TYFCB/TYDCB Report sheet")
            return {}
        
        print(f"📋 Reading TYFCB from sheet: {found_sheet_name}")
        
        data = {}
        all_rows = list(tyfcb_sheet.iter_rows(values_only=True))
        
        if len(all_rows) < 2:
            print("❌ TYFCB sheet structure is invalid")
            return {}
        
        # Expected columns: Member, Total OTO, Unique Partners, Referrals Given, Referrals Received, etc.
        # Row 0: Title or Headers
        # Rows 1+: Data
        
        # Find header row (contains "member" or "name")
        header_idx = 0
        for idx, row in enumerate(all_rows[:3]):
            if row and any('member' in str(col).lower() or 'name' in str(col).lower() for col in row if col):
                header_idx = idx
                break
        
        # Map column names to indices
        header_row = all_rows[header_idx]
        col_map = {}
        for idx, col_name in enumerate(header_row):
            if col_name:
                col_lower = str(col_name).lower()
                if 'member' in col_lower or 'name' in col_lower:
                    col_map['member'] = idx
                elif 'oto' in col_lower or 'meetings' in col_lower:
                    col_map['oto'] = idx
                elif 'partner' in col_lower:
                    col_map['partners'] = idx
                elif 'referral' in col_lower and 'given' in col_lower:
                    col_map['ref_given'] = idx
                elif 'referral' in col_lower and 'received' in col_lower:
                    col_map['ref_received'] = idx
                elif 'inside' in col_lower:
                    col_map['inside_aed'] = idx
                elif 'outside' in col_lower:
                    col_map['outside_aed'] = idx
                elif 'total' in col_lower and 'aed' in col_lower:
                    col_map['total_aed'] = idx
                elif 'avg' in col_lower and 'referral' in col_lower:
                    col_map['avg_ref_month'] = idx
                elif 'avg' in col_lower and 'value' in col_lower:
                    col_map['avg_value'] = idx
        
        row_count = 0
        for row_idx, row_data in enumerate(all_rows[header_idx + 1:]):
            if not row_data or not row_data[0]:
                continue
            
            member = str(row_data[col_map['member']]).strip() if 'member' in col_map else None
            if not member or '__EMPTY' in member or member.lower() == 'member':
                continue
            
            # Extract scores
            scores = {'member': member}
            
            for key, col_idx in col_map.items():
                if key != 'member' and col_idx < len(row_data):
                    try:
                        val = row_data[col_idx]
                        if val and str(val).lower() not in ('nan', ''):
                            scores[key] = float(val) if '.' in str(val) else int(val)
                    except (ValueError, TypeError):
                        pass
            
            if len(scores) > 1:  # More than just member name
                data[member] = scores
                row_count += 1
        
        print(f"  ✅ Processed {row_count} rows")
        record_count = len(data)
        print(f"  ✅ Found {record_count} contribution records")
        return data
    
    except Exception as e:
        print(f"❌ Error reading TYFCB Report: {e}")
        return {}

def generate_interaction_sql(interactions):
    """Generate SQL INSERT for interactions"""
    if not interactions:
        print("❌ No interactions to import")
        return None
    
    # Build VALUES clause
    values = []
    for key, data in interactions.items():
        a = data['a'].replace("'", "''")
        b = data['b'].replace("'", "''")
        c = data['c']
        values.append(f"('{a}', '{b}', {c})")
    
    sql = f"""BEGIN;

DELETE FROM lad_dev.community_roi_interactions 
WHERE tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid;

INSERT INTO lad_dev.community_roi_interactions 
(id, tenant_id, member_a_id, member_b_id, meeting_count, metadata, created_at, updated_at, is_deleted)
SELECT 
    gen_random_uuid(),
    '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid,
    ma.id,
    mb.id,
    pairs.meeting_count,
    jsonb_build_object('source', 'import_one_to_one_matrix'),
    NOW(),
    NOW(),
    false
FROM (
    VALUES 
        {','.join(values)}
) AS pairs(member_a_name, member_b_name, meeting_count)
JOIN lad_dev.community_roi_members ma ON ma.name = pairs.member_a_name
JOIN lad_dev.community_roi_members mb ON mb.name = pairs.member_b_name
WHERE ma.tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid
  AND ma.is_deleted = false
  AND mb.tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid
  AND mb.is_deleted = false
  AND ma.id != mb.id;

-- Verify interactions
SELECT COUNT(*) as total_interactions FROM lad_dev.community_roi_interactions
WHERE tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid AND is_deleted = false;

COMMIT;
"""
    return sql

def generate_referral_sql(referrals):
    """Generate SQL INSERT for referrals"""
    if not referrals:
        print("❌ No referrals to import")
        return None
    
    # Build VALUES clause
    values = []
    for key, data in referrals.items():
        a = data['a'].replace("'", "''")
        b = data['b'].replace("'", "''")
        c = data['c']
        values.append(f"('{a}', '{b}', {c})")
    
    sql = f"""BEGIN;

DELETE FROM lad_dev.community_roi_referrals 
WHERE tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid;

INSERT INTO lad_dev.community_roi_referrals 
(id, tenant_id, referred_by_id, referred_to_id, referral_count, metadata, created_at, updated_at, is_deleted)
SELECT 
    gen_random_uuid(),
    '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid,
    referred_by.id,
    referred_to.id,
    pairs.referral_count,
    jsonb_build_object('source', 'import_referral_matrix'),
    NOW(),
    NOW(),
    false
FROM (
    VALUES 
        {','.join(values)}
) AS pairs(referred_by_name, referred_to_name, referral_count)
JOIN lad_dev.community_roi_members referred_by ON referred_by.name = pairs.referred_by_name
JOIN lad_dev.community_roi_members referred_to ON referred_to.name = pairs.referred_to_name
WHERE referred_by.tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid
  AND referred_by.is_deleted = false
  AND referred_to.tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid
  AND referred_to.is_deleted = false
  AND referred_by.id != referred_to.id;

-- Verify referrals
SELECT COUNT(*) as total_referrals FROM lad_dev.community_roi_referrals
WHERE tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid AND is_deleted = false;

COMMIT;
"""
    return sql

def generate_combination_sql(combinations):
    """Generate SQL INSERT for combination matrix"""
    if not combinations:
        print("❌ No combinations to import")
        return None
    
    # Build VALUES clause
    values = []
    for key, data in combinations.items():
        a = data['a'].replace("'", "''")
        b = data['b'].replace("'", "''")
        t = data['t']
        values.append(f"('{a}', '{b}', {t})")
    
    sql = f"""BEGIN;

DELETE FROM lad_dev.community_roi_relationship_scores 
WHERE tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid;

INSERT INTO lad_dev.community_roi_relationship_scores 
(id, tenant_id, member_a_id, member_b_id, combination_type, metadata, created_at, updated_at, is_deleted)
SELECT 
    gen_random_uuid(),
    '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid,
    ma.id,
    mb.id,
    pairs.combo_type,
    jsonb_build_object('source', 'import_combination_matrix'),
    NOW(),
    NOW(),
    false
FROM (
    VALUES 
        {','.join(values)}
) AS pairs(member_a_name, member_b_name, combo_type)
JOIN lad_dev.community_roi_members ma ON ma.name = pairs.member_a_name
JOIN lad_dev.community_roi_members mb ON mb.name = pairs.member_b_name
WHERE ma.tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid
  AND ma.is_deleted = false
  AND mb.tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid
  AND mb.is_deleted = false
  AND ma.id != mb.id;

-- Verify combinations
SELECT COUNT(*) as total_combinations FROM lad_dev.community_roi_relationship_scores
WHERE tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid AND is_deleted = false;

COMMIT;
"""
    return sql

def generate_tyfcb_sql(tyfcb_data):
    """Generate SQL INSERT for TYFCB Report -> contribution_scores"""
    if not tyfcb_data:
        print("❌ No TYFCB data to import")
        return None
    
    # Build VALUES clause
    values = []
    for member_name, scores in tyfcb_data.items():
        name = member_name.replace("'", "''")
        total_oto = scores.get('oto', 0)
        partners = scores.get('partners', 0)
        ref_given = scores.get('ref_given', 0)
        ref_received = scores.get('ref_received', 0)
        inside_aed = scores.get('inside_aed', 0)
        outside_aed = scores.get('outside_aed', 0)
        total_aed = scores.get('total_aed', 0)
        avg_ref_month = scores.get('avg_ref_month', 0)
        avg_value = scores.get('avg_value', 0)
        
        values.append(f"('{name}', {total_oto}, {partners}, {ref_given}, {ref_received}, {inside_aed}, {outside_aed}, {total_aed}, {avg_ref_month}, {avg_value})")
    
    sql = f"""BEGIN;

DELETE FROM lad_dev.community_roi_contribution_scores 
WHERE tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid;

INSERT INTO lad_dev.community_roi_contribution_scores 
(id, tenant_id, member_id, total_oto, unique_oto_partners, total_referrals_given, total_referrals_received, 
 inside_network_business_aed, outside_network_business_aed, total_business_aed, 
 avg_referrals_per_month, avg_value_per_referral_aed, calculated_at, metadata, created_at, updated_at, is_deleted)
SELECT 
    gen_random_uuid(),
    '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid,
    m.id,
    scores.total_oto,
    scores.partners,
    scores.ref_given,
    scores.ref_received,
    scores.inside_aed,
    scores.outside_aed,
    scores.total_aed,
    scores.avg_ref_month,
    scores.avg_value,
    NOW(),
    jsonb_build_object('source', 'import_tyfcb_report'),
    NOW(),
    NOW(),
    false
FROM (
    VALUES 
        {','.join(values)}
) AS scores(member_name, total_oto, partners, ref_given, ref_received, inside_aed, outside_aed, total_aed, avg_ref_month, avg_value)
JOIN lad_dev.community_roi_members m ON m.name = scores.member_name
WHERE m.tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid
  AND m.is_deleted = false;

-- Verify contributions
SELECT COUNT(*) as total_contributions FROM lad_dev.community_roi_contribution_scores
WHERE tenant_id = '9ca4012a-2e02-5593-8cc1-fd5bd81483f9'::uuid AND is_deleted = false;

COMMIT;
"""
    return sql

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python3 extract_all_sheets.py <excel_file> [sheet_type]")
        print("\nAvailable sheet types: all, interactions, referrals, combination, tyfcb")
        print("  all          - Extract all four sheet types (default)")
        print("  interactions - One-to-One Matrix only")
        print("  referrals    - Referral Matrix only")
        print("  combination  - Combination Matrix only")
        print("  tyfcb        - TYFCB Report only")
        sys.exit(1)
    
    excel_file = sys.argv[1]
    sheet_type = sys.argv[2].lower() if len(sys.argv) > 2 else 'all'
    
    print(f"📂 Reading: {excel_file}\n")
    
    results = {}
    
    # Extract based on sheet type
    if sheet_type in ('all', 'interactions'):
        print("🔍 Processing One-to-One Matrix...")
        interactions = read_excel_heatmap(excel_file)
        if interactions:
            results['interactions'] = interactions
    
    if sheet_type in ('all', 'referrals'):
        print("\n🔍 Processing Referral Matrix...")
        referrals = read_excel_referral_matrix(excel_file)
        if referrals:
            results['referrals'] = referrals
    
    if sheet_type in ('all', 'combination'):
        print("\n🔍 Processing Combination Matrix...")
        combinations = read_excel_combination_matrix(excel_file)
        if combinations:
            results['combination'] = combinations
    
    if sheet_type in ('all', 'tyfcb'):
        print("\n🔍 Processing TYFCB Report...")
        tyfcb = read_excel_tyfcb_report(excel_file)
        if tyfcb:
            results['tyfcb'] = tyfcb
    
    print("\n" + "="*60)
    
    if not results:
        print("❌ No data to import")
        sys.exit(1)
    
    # Generate and save SQL files
    generated_files = []
    
    if 'interactions' in results:
        sql_int = generate_interaction_sql(results['interactions'])
        if sql_int:
            with open('/tmp/import_interactions.sql', 'w') as f:
                f.write(sql_int)
            generated_files.append('interactions')
            print(f"✅ Generated: /tmp/import_interactions.sql")
    
    if 'referrals' in results:
        sql_ref = generate_referral_sql(results['referrals'])
        if sql_ref:
            with open('/tmp/import_referrals.sql', 'w') as f:
                f.write(sql_ref)
            generated_files.append('referrals')
            print(f"✅ Generated: /tmp/import_referrals.sql")
    
    if 'combination' in results:
        sql_comb = generate_combination_sql(results['combination'])
        if sql_comb:
            with open('/tmp/import_combination.sql', 'w') as f:
                f.write(sql_comb)
            generated_files.append('combination')
            print(f"✅ Generated: /tmp/import_combination.sql")
    
    if 'tyfcb' in results:
        sql_tyf = generate_tyfcb_sql(results['tyfcb'])
        if sql_tyf:
            with open('/tmp/import_tyfcb.sql', 'w') as f:
                f.write(sql_tyf)
            generated_files.append('tyfcb')
            print(f"✅ Generated: /tmp/import_tyfcb.sql")
    
    print(f"\n📌 Generated files summary:")
    print(json.dumps({k: len(v) for k, v in results.items()}, indent=2))
    
    print(f"\n📋 To import via API:")
    for file_type in generated_files:
        print(f"  - {file_type.upper()}")
